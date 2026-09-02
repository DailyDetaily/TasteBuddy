#!/usr/bin/env python3
"""Extract Korean food name/category metadata without nutrient values."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


LATEST_DB_SHEET = "국가표준식품성분 Database 10.4"
APPENDIX_NAME_PREFIX = "부록2)"


def compact(value: Any) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def stable_id(prefix: str, value: str) -> str:
    digest = hashlib.sha1(value.encode("utf-8")).hexdigest()[:12]
    return f"{prefix}:{digest}"


def split_aliases(ko_name: str, english_name: str) -> list[str]:
    aliases: list[str] = []

    def push(value: str) -> None:
        value = compact(value)
        if value and value not in aliases:
            aliases.append(value)

    push(ko_name)

    korean_parts = [compact(part) for part in ko_name.split(",")]
    if korean_parts:
        push(korean_parts[0])
    for part in korean_parts[1:]:
        if len(part) >= 2 and not any(token in part for token in ["생것", "익힌것", "말린것"]):
            push(part)

    english_base = english_name.split("(")[0].split(",")[0].strip()
    push(english_base)
    push(english_name)

    return aliases[:10]


def build_appendix_index(workbook) -> dict[str, dict[str, str]]:
    appendix_sheet = next(
        name for name in workbook.sheetnames if name.startswith(APPENDIX_NAME_PREFIX)
    )
    sheet = workbook[appendix_sheet]
    rows = sheet.iter_rows(min_row=2, values_only=True)
    index: dict[str, dict[str, str]] = {}

    for row in rows:
        db_index = compact(row[0])
        if not db_index:
            continue

        index[db_index] = {
            "dbIndex": db_index,
            "foodCode": compact(row[1]),
            "koName": compact(row[2]),
            "englishName": compact(row[3]),
            "scientificName": compact(row[4]),
        }

    return index


def extract_catalog(source_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    appendix = build_appendix_index(workbook)
    sheet = workbook[LATEST_DB_SHEET]
    entries: list[dict[str, Any]] = []
    seen: set[str] = set()

    for row in sheet.iter_rows(min_row=4, values_only=True):
        db_index = compact(row[0])
        food_group = compact(row[2])
        db_ko_name = compact(row[3])

        if not db_index or not food_group or not db_ko_name:
            continue

        appendix_item = appendix.get(db_index, {})
        ko_name = appendix_item.get("koName") or db_ko_name
        english_name = appendix_item.get("englishName", "")
        scientific_name = appendix_item.get("scientificName", "")
        food_code = appendix_item.get("foodCode", "")
        dedupe_key = "|".join([food_group, ko_name, english_name, scientific_name])

        if dedupe_key in seen:
            continue

        seen.add(dedupe_key)
        source_key = food_code or f"db-index:{db_index}"

        entries.append({
            "id": stable_id("korean-food", source_key),
            "dbIndex": db_index,
            "foodCode": food_code,
            "koName": ko_name,
            "englishName": english_name,
            "scientificName": scientific_name,
            "foodGroup": food_group,
            "aliases": split_aliases(ko_name, english_name),
            "source": "korean-standard-food-composition-db",
            "sourceVersion": "10.4",
        })

    entries.sort(key=lambda item: (item["foodGroup"], item["koName"]))
    return entries


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    source_path = Path(args.source)
    out_path = Path(args.out)
    catalog = extract_catalog(source_path)
    payload = {
        "version": "0.1",
        "source": "korean-standard-food-composition-db",
        "sourcePath": str(source_path),
        "count": len(catalog),
        "items": catalog,
    }

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Saved {len(catalog)} Korean standard food entries to {out_path}")


if __name__ == "__main__":
    main()
